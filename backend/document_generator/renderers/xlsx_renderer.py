# Excel Spreadsheet (.xlsx) Strategy Renderer
from __future__ import annotations

import os
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from .base import Renderer
from ..schemas import XlsxSpec, DocumentResult, DocType


class XlsxRenderer(Renderer):

    def render(self, spec: XlsxSpec, output_path: Path) -> DocumentResult:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        # Excel forbids characters \ / ? * : [ ] in sheet names and caps at 31 characters
        clean_sheet_name = re.sub(r"[\\/*?:\[\]]", "_", (spec.sheet_name or "Sheet1").strip())[:31] or "Sheet1"
        ws.title = clean_sheet_name

        header_hex = "{:02X}{:02X}{:02X}".format(*self.theme.rgb_dark)
        for col_num, h in enumerate(spec.headers, 1):
            cell = ws.cell(row=1, column=col_num, value=h)
            cell.font = Font(name=self.theme.font_family, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=header_hex, end_color=header_hex, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, row in enumerate(spec.rows, 2):
            for c_idx, val in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = Font(name=self.theme.font_family)
                self._apply_rules(cell, val, spec.rules)

        # Optimize column widths for Windows Excel display (preventing '###' and truncated text)
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(str(output_path))

        return DocumentResult(
            filename=output_path.name,
            path=str(output_path),
            doc_type=DocType.XLSX,
            size_bytes=os.path.getsize(output_path),
        )

    def _apply_rules(self, cell, val, rules) -> None:
        """Applies solid cell fill highlighting dynamically based on rule.color_hex."""
        val_str = str(val).strip().upper()
        for rule in rules:
            if any(m.upper() == val_str or m.upper() in val_str.split() for m in rule.match_values):
                color = str(rule.color_hex or "{:02X}{:02X}{:02X}".format(*self.theme.rgb_orange)).lstrip("#").upper()
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                cell.font = Font(name=self.theme.font_family, bold=rule.bold, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center")
                return
